"""Build index.html from the Garden AIO workbook.

Usage:  python build/build_app.py [path/to/Garden_AIO_Data.xlsx]
Reads the recalculated workbook (open and save it in Excel first so formula
values are cached), exports the data the app needs, and writes ../index.html.
Requires: pip install openpyxl
"""
import json, sys, pathlib
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
xlsx = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "Garden_AIO_Data.xlsx"
wb = load_workbook(xlsx, data_only=True)
cr, co, en, cl, vt = wb["2_Crops"], wb["3_Companions"], wb["2b_Environment"], wb["10_Climate"], wb["2c_Varieties"]

PET = {"Safe": "S", "Caution": "C", "Toxic": "T", "Verify": "V"}
EVID = {"Well-supported": "W", "Traditional": "T", "Uncertain": "U"}
env_row = {en.cell(row=r, column=1).value: r for r in range(5, 320) if en.cell(row=r, column=1).value}

plants = []
for r in range(5, 320):
    n = cr.cell(row=r, column=1).value
    if not n or cr.cell(row=r, column=36).value != "Y":      # only rows marked "Bed designer ready = Y"
        continue
    g = lambda c: cr.cell(row=r, column=c).value
    e = env_row[n]
    ev = lambda c: en.cell(row=e, column=c).value
    months = [int(x) for x in str(g(37)).split(",")] if g(37) else []
    plants.append(dict(n=n, ty=g(2), s=g(5), sfg=g(6), h=g(12), pet=PET[g(16)], pn=g(17), win=g(9), w=months,
                       sp=g(28), lo=g(29), hi=g(30), ss=g(31), dm=g(32), di=g(33), dsrc=g(34),
                       dl=ev(17), do=ev(18), h1=ev(20), h2=ev(21)))
ready = {p["n"] for p in plants}

def mid(a, b):
    if a is None and b is None: return None
    a2 = a if a is not None else b; b2 = b if b is not None else a
    return round((a2 + b2) / 2)

def trim(t, n=140):
    return t if len(t) <= n else t[:n].rsplit(" ", 1)[0] + "..."

varieties = {}
for r in range(5, 2505):
    ws, name = vt.cell(row=r, column=22).value, vt.cell(row=r, column=3).value
    if not ws or ws not in ready or not name: continue
    g = lambda c: vt.cell(row=r, column=c).value
    lo, hi, hmin, hmax = g(16), g(17), g(18), g(19)
    src = g(23) or ""
    src = "Wind River Greens (CC BY 4.0)" if src.startswith("Wind River") else src
    varieties.setdefault(ws, []).append([name, g(4), mid(lo, hi), lo, hi, mid(hmin, hmax), g(15), g(21) or "", trim(g(8) or ""), src, g(24)])
for k in varieties: varieties[k].sort(key=lambda v: v[0].lower())

mech, midx, pairs = [], {}, []
for r in range(12, 700):
    a = co.cell(row=r, column=1).value
    if not a: break
    b, rel, sc, m, evd = [co.cell(row=r, column=c).value for c in (2, 3, 4, 5, 6)]
    if m not in midx: midx[m] = len(mech); mech.append(m)
    pairs.append([a, b, int(sc), EVID[evd], midx[m]])

solar = [cl.cell(row=11 + i, column=3).value for i in range(12)]
data = dict(plants=plants, pairs=pairs, mech=mech, solar=solar, vars=varieties)

# optional photos: images.json at the repo root (see images.example.json). Your entries replace the
# Wikipedia photos the app loads at runtime.
img_path = ROOT / "images.json"
data["img"] = json.loads(img_path.read_text(encoding="utf-8")) if img_path.exists() else {}

# constants for the Materials calculator (sources are listed in the workbook tab 13_Materials)
const = {"mix": [["UMN: 60% topsoil, 40% compost", .6, .4, 0, "", ""],
                 ["Iowa State: 70% topsoil, 30% compost", .7, .3, 0, "", ""],
                 ["Iowa State: equal thirds topsoil, compost, coarse sand", .334, .333, .333, "Coarse sand", "sand"],
                 ["Maryland: compost and topsoil 1:2", .667, .333, 0, "", ""],
                 ["Maryland: compost and soilless mix 1:1", 0, .5, .5, "Soilless mix", "soilless"]],
         "dens": {"topsoil": 2200, "compost": 1250, "sand": 2600, "soilless": 600, "mulch": 600, "gravel": 2700},
         "wet": 1.3, "gal": 0.623, "topdress": 2}

html = (ROOT / "build" / "app_template.html").read_text(encoding="utf-8")
html = html.replace("/*DATA*/", json.dumps(data, separators=(",", ":"))).replace("/*CONST*/", json.dumps(const, separators=(",", ":")))
(ROOT / "index.html").write_text(html, encoding="utf-8")
print(f"index.html written: {len(plants)} plants, {len(pairs)} pairs, {sum(len(v) for v in varieties.values())} varieties, "
      f"{len(data['img'].get('plants', {}))} custom plant photos, {len(data['img'].get('varieties', {}))} custom variety photos")
